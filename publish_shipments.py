#!/usr/bin/env python3
"""
publish_shipments.py -- 出船管理（3社共用出貨排程）同步腳本

讀取「銘環船_出船管理」Excel，把逐筆訂單的 HTW 検査排程／出貨資訊、
以及船期參考表，同步進 Supabase（shipments / sailing_schedule 兩張表）。

跟 publish.py 是同樣的操作習慣，但資料來源、schema 完全獨立：
  - shipments：只做 upsert，不會自動刪除 Excel 裡消失的舊列
    （避免連動刪掉該筆訂單底下的留言討論歷史）
  - sailing_schedule：純參考資料、無留言，每次同步用「全刪重建」

使用方式：
    python publish_shipments.py            # 正式執行
    python publish_shipments.py --dry-run  # 只顯示解析結果，不連線 Supabase

環境變數（可寫在同目錄的 .env，或用系統環境變數）：
    SUPABASE_URL               同 publish.py
    SUPABASE_SERVICE_ROLE_KEY  同 publish.py
    SHIPMENTS_EXCEL_PATH       預設見下方 DEFAULT_EXCEL_PATH
"""
import argparse
import datetime
import os
import sys
from pathlib import Path

import requests

from publish import load_dotenv, SupabaseClient

DEFAULT_EXCEL_PATH = r"C:\Users\q5695\Desktop\【送信用】銘環船_出船管理.xlsx"

SHEET1_DATA_START_ROW = 3
SHEET2_DATA_START_ROW = 5

EXCEL_EPOCH = datetime.date(1899, 12, 30)


def parse_yyyymmdd(value):
    """発注納期のような "20260810" 形式（8桁の数字／文字列）を date に変換する。"""
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    s = str(value).strip()
    if not s:
        return None
    try:
        return datetime.datetime.strptime(s, "%Y%m%d").date()
    except ValueError:
        return None


def parse_excel_serial(value):
    """HTW検査完了日のような Excel シリアル値（例："46230"）を date に変換する。
    セルが最初から日付書式で datetime になっている場合はそのまま使う。
    人が手入力した "2026.09.30" 等のテキスト日付（Excel 側で日付として
    認識されていないケース）も、フォールバックとしてパースを試みる。"""
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    s = str(value).strip()
    if not s:
        return None
    try:
        serial = float(s)
        if serial <= 0:
            return None
        return EXCEL_EPOCH + datetime.timedelta(days=serial)
    except (TypeError, ValueError):
        pass
    for fmt in ("%Y.%m.%d", "%Y/%m/%d", "%Y-%m-%d"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_number(value):
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def clean_text(value):
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def extract_shipments(ws):
    rows = []
    skipped = []
    for r in range(SHEET1_DATA_START_ROW, ws.max_row + 1):
        order_no = clean_text(ws.cell(row=r, column=4).value)
        line_no = clean_text(ws.cell(row=r, column=5).value)
        customer = clean_text(ws.cell(row=r, column=2).value)

        if not order_no and not line_no and not customer:
            continue  # 完全空白列，略過
        if not order_no or not line_no:
            skipped.append((r, "缺少 発注No 或 行No，無法組成唯一 id"))
            continue

        rows.append({
            "id": f"{order_no}_{line_no}",
            "row_order": r,
            "customer": customer,
            "part_no": clean_text(ws.cell(row=r, column=3).value),
            "order_no": order_no,
            "line_no": line_no,
            "order_qty": parse_number(ws.cell(row=r, column=6).value),
            "delivery_type": clean_text(ws.cell(row=r, column=7).value),
            "order_due_date": iso(parse_yyyymmdd(ws.cell(row=r, column=8).value)),
            "htw_delivery_date": iso(parse_excel_serial(ws.cell(row=r, column=9).value)),
            "htw_inspection_planned_date": iso(parse_excel_serial(ws.cell(row=r, column=10).value)),
            "htw_inspection_minutes": parse_number(ws.cell(row=r, column=11).value),
            "htw_inspection_done_date": iso(parse_excel_serial(ws.cell(row=r, column=12).value)),
            "mh_pickup_date": iso(parse_excel_serial(ws.cell(row=r, column=13).value)),
            "twh_ship_month": clean_text(ws.cell(row=r, column=14).value),
            "twh_ship_vessel": clean_text(ws.cell(row=r, column=15).value),
            "komaki_ship_month": clean_text(ws.cell(row=r, column=16).value),
            "komaki_ship_vessel": clean_text(ws.cell(row=r, column=17).value),
            "japan_arrival_date": iso(parse_excel_serial(ws.cell(row=r, column=18).value)),
            "excel_comment": clean_text(ws.cell(row=r, column=19).value),
        })
    return rows, skipped


def extract_sailing_schedule(ws):
    rows = []
    for r in range(SHEET2_DATA_START_ROW, ws.max_row + 1):
        seq_no = ws.cell(row=r, column=12).value  # L
        if seq_no is None:
            continue
        rows.append({
            "seq_no": int(seq_no) if str(seq_no).strip().isdigit() else None,
            "meikan_pickup_range": clean_text(ws.cell(row=r, column=13).value),   # M
            "htw_taichung_arrival": clean_text(ws.cell(row=r, column=14).value),  # N
            "schedule_label": clean_text(ws.cell(row=r, column=15).value),        # O
            "tw_closing_date": clean_text(ws.cell(row=r, column=16).value),       # P
            "jp_eta_date": iso(parse_excel_serial(ws.cell(row=r, column=17).value)),   # Q
            "komaki_arrival_date": iso(parse_excel_serial(ws.cell(row=r, column=18).value)),  # R
            "note": clean_text(ws.cell(row=r, column=19).value),                  # S
        })
    return rows


def iso(d):
    return d.isoformat() if d else None


def main():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")

    parser = argparse.ArgumentParser(description="同步出船管理排程到 Supabase")
    parser.add_argument("--excel", default=None, help="Excel 檔案路徑（覆蓋預設值）")
    parser.add_argument("--dry-run", action="store_true", help="只顯示解析結果，不連線 Supabase")
    args = parser.parse_args()

    script_dir = Path(__file__).resolve().parent
    load_dotenv(script_dir / ".env")

    excel_path = Path(args.excel or os.environ.get("SHIPMENTS_EXCEL_PATH", DEFAULT_EXCEL_PATH))
    if not excel_path.exists():
        sys.exit(f"[錯誤] 找不到 Excel 檔案：{excel_path}")

    import openpyxl
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws1 = wb.worksheets[0]
    ws2 = wb.worksheets[1]

    print(f"讀取 Excel：{excel_path}")
    shipments, skipped = extract_shipments(ws1)
    schedule = extract_sailing_schedule(ws2)
    print(f"出船排程：{len(shipments)} 筆，船期參考：{len(schedule)} 筆")
    if skipped:
        print(f"略過：{len(skipped)} 列")
        for r, reason in skipped:
            print(f"  - Excel row {r}：{reason}")

    if args.dry_run:
        print("\n--dry-run 模式，以下為預覽（節錄前 10 筆與後 5 筆），不會寫入 Supabase：\n")
        preview_rows = shipments[:10] + (shipments[-5:] if len(shipments) > 10 else [])
        for row in preview_rows:
            print(f"[{row['id']}] {row['customer']} / {row['part_no']} / "
                  f"発注納期={row['order_due_date']} HTW完了={row['htw_inspection_done_date']} "
                  f"TWH船={row['twh_ship_month']}{row['twh_ship_vessel'] or ''} "
                  f"小牧船={row['komaki_ship_month']}{row['komaki_ship_vessel'] or ''} "
                  f"日本入荷={row['japan_arrival_date']}")
        print("\n船期參考表節錄：")
        for row in schedule[:5]:
            print(f"  No.{row['seq_no']} {row['schedule_label']} "
                  f"ETA={row['jp_eta_date']} 小牧着={row['komaki_arrival_date']}")
        return

    supabase_url = os.environ.get("SUPABASE_URL")
    service_key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if not supabase_url or not service_key:
        sys.exit(
            "[錯誤] 缺少 SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY。\n"
            "請在 publish_shipments.py 同目錄建立 .env（可參考 .env.example）或設定環境變數。"
        )

    client = SupabaseClient(supabase_url, service_key)
    headers = {**client.headers, "Content-Type": "application/json"}

    # shipments：只 upsert，不刪除
    upserted = 0
    for row in shipments:
        resp = requests.post(
            f"{client.url}/rest/v1/shipments",
            headers={**headers, "Prefer": "resolution=merge-duplicates,return=minimal"},
            params={"on_conflict": "id"},
            json={**row, "updated_at": "now()"},
            timeout=30,
        )
        resp.raise_for_status()
        upserted += 1
    print(f"shipments upsert：{upserted} 筆")

    # sailing_schedule：全刪重建
    resp = requests.delete(
        f"{client.url}/rest/v1/sailing_schedule",
        headers=headers,
        params={"id": "gte.0"},
        timeout=30,
    )
    resp.raise_for_status()
    if schedule:
        resp = requests.post(
            f"{client.url}/rest/v1/sailing_schedule",
            headers={**headers, "Prefer": "return=minimal"},
            json=schedule,
            timeout=30,
        )
        resp.raise_for_status()
    print(f"sailing_schedule 全刪重建：{len(schedule)} 筆")


if __name__ == "__main__":
    main()
